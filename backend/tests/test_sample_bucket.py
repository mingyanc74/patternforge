import csv
import io
import json
import uuid
from datetime import timedelta
from unittest.mock import Mock

import pytest
from core.errors import PipelineError
from core.models import Job, S3Connection
from core.sample_bucket import BUCKET, CATALOG, SampleBucket, object_bytes
from core.services.s3 import client_for, fingerprint
from core.tasks import run_spark
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

pytestmark = pytest.mark.django_db


def sample_connection(client, settings):
    settings.SAMPLE_BUCKET_ENABLED = True
    response = client.post("/api/s3/sample-connection", {}, format="json")
    assert response.status_code == 201, response.data
    return S3Connection.objects.get(pk=response.data["id"])


def test_sample_is_explicit_optional_and_parameter_free(client, settings):
    settings.SAMPLE_BUCKET_ENABLED = False
    assert client.get("/api/bootstrap").data["sample_bucket_enabled"] is False
    assert client.post("/api/s3/sample-connection", {}, format="json").status_code == 404
    settings.SAMPLE_BUCKET_ENABLED = True
    for body in [{"bucket": "arbitrary"}, {"endpoint": "http://example.com"}, {"access_key": "secret"}]:
        assert client.post("/api/s3/sample-connection", body, format="json").status_code == 400
    item = sample_connection(client, settings)
    assert item.source_type == "sample" and item.ciphertext == "" and item.bucket == BUCKET
    again = client.post("/api/s3/sample-connection", {}, format="json")
    assert again.status_code == 200 and again.data["id"] == str(item.id)
    assert len(client.get("/api/s3/connections").data["connections"]) == 1
    assert "secret" not in again.data and "ciphertext" not in again.data


def test_sample_has_csrf_and_session_isolation(client, settings):
    item = sample_connection(client, settings)
    strict = APIClient(enforce_csrf_checks=True)
    boot = strict.get("/api/bootstrap")
    assert strict.post("/api/s3/sample-connection", {}, format="json").status_code == 403
    response = strict.post(
        "/api/s3/sample-connection", {}, format="json", HTTP_X_CSRFTOKEN=boot.data["csrf_token"]
    )
    assert response.status_code == 201 and response.data["id"] != str(item.id)
    assert strict.get(f"/api/s3/connections/{item.id}/files").status_code == 404
    assert strict.delete(f"/api/s3/connections/{item.id}").status_code == 403
    assert strict.get("/api/jobs").data["jobs"] == []


def test_sample_reads_use_no_aws_credentials_or_network(client, settings, mocker):
    aws = mocker.patch("core.services.s3.boto3.client", side_effect=AssertionError("AWS must not be used"))
    item = sample_connection(client, settings)
    files = client.get(f"/api/s3/connections/{item.id}/files")
    assert files.status_code == 200 and len(files.data["files"]) == len(CATALOG)
    assert all(entry["description"] for entry in files.data["files"])
    assert [entry for entry in files.data["files"] if entry["key"].endswith(".txt")][0]["supported"] is False
    assert fingerprint(client_for(item), BUCKET, "examples/customers.csv")["size"] > 10000
    response = client.post(
        "/api/files/inspect",
        {
            "connection_id": str(item.id),
            "source_key": "examples/customers.csv",
        },
        format="json",
    )
    assert response.status_code == 202 and response.data["source_type"] == "sample"
    assert response.data["status"] == "QUEUED"
    aws.assert_not_called()


def test_sample_only_catalog_keys_and_rejects_invalid_files(client, settings):
    item = sample_connection(client, settings)
    for key, code in [
        ("../.env", "SAMPLE_NOT_FOUND"),
        ("edge-cases/empty.csv", "EMPTY_FILE"),
        ("edge-cases/unsupported.txt", "UNSUPPORTED_FILE"),
    ]:
        response = client.post(
            "/api/files/inspect",
            {
                "connection_id": str(item.id),
                "source_key": key,
            },
            format="json",
        )
        assert response.status_code == 400 and response.data["error"]["code"] == code
    assert Job.objects.count() == 0
    listing = client.get(f"/api/s3/connections/{item.id}/files", {"prefix": "examples/"})
    assert len(listing.data["files"]) == 3


def test_sample_expiry_disconnect_and_reconnect(client, settings):
    item = sample_connection(client, settings)
    client.delete(f"/api/s3/connections/{item.id}")
    assert client.get(f"/api/s3/connections/{item.id}/files").status_code == 410
    new = sample_connection(client, settings)
    assert new.id != item.id
    new.expires_at = timezone.now() - timedelta(seconds=1)
    new.save()
    assert client.get(f"/api/s3/connections/{new.id}/files").status_code == 410
    assert sample_connection(client, settings).id != new.id


def test_real_s3_request_cannot_switch_provider(client, bucket, settings):
    settings.SAMPLE_BUCKET_ENABLED = True
    response = client.post(
        "/api/s3/connections",
        {
            "bucket": "user-owned-bucket",
            "access_key": "test-user",
            "secret_key": "test-secret",
            "source_type": "sample",
        },
        format="json",
    )
    assert response.status_code == 201 and response.data["source_type"] == "amazon_s3"
    assert client.post("/api/s3/connections", {}, format="json").status_code == 400


def test_synthetic_data_is_stable_and_covers_csv_excel_edges():
    rows = list(csv.DictReader(io.StringIO(object_bytes("examples/customers.csv").decode())))
    assert len(rows) == 1000 and rows[0]["ID"] == "000001"
    assert rows[2]["Name"] == "Bob, Taylor" and rows[3]["Name"] == "Two\nLines"
    assert sum("@" in row["Email"] for row in rows) == 800
    assert sum(row["Email"] == "\\N" for row in rows) == 100
    key = "examples/workbook.xlsx"
    first = object_bytes(key)
    object_bytes.cache_clear()
    assert first == object_bytes(key)
    book = load_workbook(io.BytesIO(first), read_only=True)
    assert book.sheetnames == ["Customers", "Orders"]
    assert book["Customers"].max_row == 121 and book["Orders"].max_row == 81
    assert book["Customers"]["A2"].value == "000001"
    book.close()
    assert len(list(csv.reader(io.StringIO(object_bytes("examples/large-customers.csv").decode())))) == 50001
    with pytest.raises(PipelineError):
        SampleBucket().head_object(Bucket="another-bucket", Key=key)


def test_sample_stages_real_spark_input_without_secret_and_cleans_it(client, settings, tmp_path, mocker):
    # Import before patching the shared subprocess module: PySpark captures
    # Popen at import time, including for its Java gateway's atexit cleanup.
    import pyspark.java_gateway  # noqa: F401

    item = sample_connection(client, settings)
    job = Job.objects.create(
        owner=item.owner, connection=item, kind="inspect", source_key="examples/customers.csv"
    )
    settings.SPARK_MASTER = "local[2]"
    lease = uuid.uuid4()

    def launch(*args, **kwargs):
        spec = json.loads((tmp_path / "spec.json").read_text())
        assert spec["lease"] == str(lease) and spec["encrypted_credentials"] == ""
        assert spec["local_source"] == str(tmp_path / "sample-input.csv")
        assert (tmp_path / "sample-input.csv").read_bytes() == object_bytes(job.source_key)
        (tmp_path / "output.json").write_text('{"preview": []}')
        return Mock(poll=Mock(return_value=0), returncode=0)

    mocker.patch("core.tasks.subprocess.Popen", side_effect=launch)
    assert run_spark(job, lease, tmp_path) == {"preview": []}
    assert not (tmp_path / "sample-input.csv").exists() and not (tmp_path / "spec.json").exists()
