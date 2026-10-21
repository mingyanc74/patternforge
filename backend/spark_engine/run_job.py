"""Isolated Spark driver; specs contain encrypted credentials, never command-line AWS secrets."""

import codecs
import csv
import io
import json
import os
import sys
import threading
import time
import zipfile
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


def atomic_json(path, value):
    temporary = Path(str(path) + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, default=str), encoding="utf-8")
    os.replace(temporary, path)


def run(spec_path):
    from cryptography.fernet import Fernet
    from openpyxl import load_workbook
    from pyspark import StorageLevel
    from pyspark.sql import SparkSession
    from pyspark.sql import functions as F
    from pyspark.sql.types import StringType, StructField, StructType

    from spark_engine.transformations import column, statistics, transform

    spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
    work = Path(spec_path).parent
    started = time.monotonic()

    def stage(name, progress):
        atomic_json(work / "progress.json", {"stage": name, "progress": progress})

    def fail(code, message):
        atomic_json(work / "failure.json", {"code": code, "message": message})
        raise RuntimeError(code)

    stage("READING", 35)
    config = spec["config"]
    builder = (
        SparkSession.builder.appName("patternforge-" + spec["job_id"])
        .config("spark.sql.shuffle.partitions", str(spec["partitions"]))
        .config("spark.redaction.regex", "(?i)secret|password|token|access.?key|credential|authorization")
        .config("spark.redaction.string.regex", "(?i)(?:AKIA|ASIA)[A-Z0-9]{16}")
        .config("spark.sql.ui.explainMode", "simple")
        .config("spark.ui.enabled", "false")
    )
    local_source = spec.get("local_source")
    credentials = None
    client = None
    if not local_source or spec.get("lease"):
        def lease_guard():
            # Independent guard: even if the Celery parent dies, an orphaned Spark driver stops.
            import django

            os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
            django.setup()
            from core.models import Job
            from django.db import close_old_connections

            last_ok = time.monotonic()
            while True:
                if time.monotonic() - started > spec["timeout"]:
                    os._exit(124)
                try:
                    close_old_connections()
                    current = Job.objects.only("lease", "status", "cancel_requested_at").get(
                        id=spec["job_id"]
                    )
                    if (
                        str(current.lease) != spec["lease"]
                        or current.status != "RUNNING"
                        or current.cancel_requested_at
                    ):
                        os._exit(130)
                    last_ok = time.monotonic()
                except Exception:
                    if time.monotonic() - last_ok > 30:
                        os._exit(75)
                time.sleep(3)

        threading.Thread(target=lease_guard, name="lease-guard", daemon=True).start()

    if local_source:
        source = str(Path(local_source).resolve())
    else:
        import boto3
        from botocore.config import Config

        credentials = json.loads(
            Fernet(os.environ["CREDENTIAL_KEY"].encode()).decrypt(spec["encrypted_credentials"].encode())
        )
        client = boto3.client(
            "s3",
            region_name=spec["region"],
            aws_access_key_id=credentials["access_key"],
            aws_secret_access_key=credentials["secret_key"],
            aws_session_token=credentials.get("session_token") or None,
            endpoint_url=spec.get("endpoint") or None,
            config=Config(connect_timeout=5, read_timeout=20, retries={"max_attempts": 2}),
        )
        source = "s3a://" + spec["bucket"] + "/" + spec["source_key"]
        builder = (
            builder.config("spark.hadoop.fs.s3a.access.key", credentials["access_key"])
            .config("spark.hadoop.fs.s3a.secret.key", credentials["secret_key"])
            .config("spark.hadoop.fs.s3a.impl", "org.apache.hadoop.fs.s3a.S3AFileSystem")
            .config("spark.hadoop.fs.s3a.endpoint.region", spec["region"])
        )
        if credentials.get("session_token"):
            builder = builder.config(
                "spark.hadoop.fs.s3a.session.token", credentials["session_token"]
            ).config(
                "spark.hadoop.fs.s3a.aws.credentials.provider",
                "org.apache.hadoop.fs.s3a.TemporaryAWSCredentialsProvider",
            )
        else:
            builder = builder.config(
                "spark.hadoop.fs.s3a.aws.credentials.provider",
                "org.apache.hadoop.fs.s3a.SimpleAWSCredentialsProvider",
            )
        if spec.get("endpoint"):
            builder = (
                builder.config("spark.hadoop.fs.s3a.endpoint", spec["endpoint"])
                .config("spark.hadoop.fs.s3a.path.style.access", "true")
                .config(
                    "spark.hadoop.fs.s3a.connection.ssl.enabled",
                    str(spec["endpoint"].startswith("https")).lower(),
                )
            )

    def read_header(path):
        if local_source or path != source:
            with open(path, "rb") as stream:
                block = stream.read(65536)
        else:
            response = client.get_object(Bucket=spec["bucket"], Key=spec["source_key"], Range="bytes=0-65535")
            with response["Body"] as stream:
                block = stream.read(65536)
        encoding = config.get("encoding", "UTF-8")
        text = codecs.getincrementaldecoder(encoding)().decode(block, final=False).lstrip("\ufeff")
        if len(block) == 65536 and "\n" not in text:
            fail("HEADER_TOO_LARGE", "The CSV header exceeds 64 KiB.")
        try:
            headers = next(csv.reader(io.StringIO(text), delimiter=config.get("delimiter", ","), strict=True))
        except StopIteration:
            fail("EMPTY_FILE", "The file contains no header or data.")
        return headers

    suffix = Path(spec["source_key"]).suffix.lower()
    if suffix == ".xlsx":
        excel_path = Path(local_source) if local_source else work / "input.xlsx"
        if not local_source:
            response = client.get_object(Bucket=spec["bucket"], Key=spec["source_key"])
            downloaded = 0
            with response["Body"] as stream, excel_path.open("wb") as target:
                while chunk := stream.read(1024 * 1024):
                    downloaded += len(chunk)
                    if downloaded > spec["max_excel_bytes"]:
                        fail("EXCEL_TOO_LARGE", "The Excel file exceeds the configured limit.")
                    target.write(chunk)
        with zipfile.ZipFile(excel_path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 1024**3:
                fail("EXCEL_EXPANSION_LIMIT", "The expanded workbook is too large.")
        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        sheets = workbook.sheetnames
        selected_sheet = config.get("sheet") or (sheets[0] if len(sheets) == 1 else None)
        if not selected_sheet:
            atomic_json(
                work / "output.json", {"sheets": sheets, "requires_sheet": True, "columns": [], "preview": []}
            )
            workbook.close()
            return
        if selected_sheet not in sheets:
            fail("SHEET_NOT_FOUND", "Select an existing worksheet.")
        rows = workbook[selected_sheet].iter_rows(values_only=True)
        first = next(rows, None)
        if first is None:
            fail("EMPTY_SHEET", "The selected sheet is empty.")
        headers = [str(value) if value is not None else "" for value in first]
        staged = work / "excel.csv"
        with staged.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            for number, row in enumerate(rows, 1):
                if number > spec["max_excel_rows"]:
                    fail("EXCEL_ROW_LIMIT", "The worksheet exceeds the row limit.")
                values = [
                    "\\N"
                    if cell is None
                    else cell.isoformat()
                    if isinstance(cell, (datetime, date))
                    else str(cell)
                    for cell in row
                ]
                if any(len(value) > spec["max_cell_chars"] for value in values):
                    fail("CELL_TOO_LARGE", "A cell exceeds the configured text limit.")
                writer.writerow(values)
        workbook.close()
        source = str(staged)
        config = {**config, "encoding": "UTF-8", "delimiter": ",", "sheet": selected_sheet}
    else:
        headers = read_header(source)
        sheets = []

    if not headers or len(headers) > 200 or any(not name.strip() or len(name) > 255 for name in headers):
        fail("INVALID_HEADER", "Use 1-200 nonempty column names of at most 255 characters.")
    if len({name.casefold() for name in headers}) != len(headers):
        fail("DUPLICATE_COLUMNS", "Column names must be unique, including case differences.")
    spark = builder.getOrCreate()
    spark.sparkContext.setLogLevel("ERROR")
    spark.sparkContext.setJobGroup(spec["job_id"], "PatternForge data processing", interruptOnCancel=True)
    try:
        schema = StructType([StructField(name, StringType(), True) for name in headers])
        df = (
            spark.read.schema(schema)
            .option("header", True)
            .option("enforceSchema", False)
            .option("mode", "FAILFAST")
            .option("multiLine", True)
            .option("quote", '"')
            .option("escape", '"')
            .option("encoding", config.get("encoding", "UTF-8"))
            .option("delimiter", config.get("delimiter", ","))
            .option("nullValue", "\\N")
            .option("emptyValue", "")
            .option("maxCharsPerColumn", spec["max_cell_chars"])
            .option("maxColumns", 200)
            .csv(source)
        )
        if spec["kind"] == "inspect":
            sample = df.limit(100)
            row_length = F.length(F.to_json(F.struct(*[column(name) for name in headers])))
            if (sample.select(row_length.alias("n")).agg(F.max("n")).first()[0] or 0) > 131072:
                fail("ROW_TOO_LARGE", "A preview row exceeds the 128 KiB safety limit.")
            preview, byte_count = [], 0
            for row in sample.collect():
                item = row.asDict()
                size = len(json.dumps(item, ensure_ascii=False).encode())
                if preview and byte_count + size > 2 * 1024**2:
                    break
                preview.append(item)
                byte_count += size
            atomic_json(
                work / "output.json",
                {
                    "sheets": sheets,
                    "sheet": config.get("sheet", ""),
                    "requires_sheet": False,
                    "columns": [{"name": name, "type": "string"} for name in headers],
                    "preview": preview,
                    "preview_limit": 100,
                    "application_id": spark.sparkContext.applicationId,
                },
            )
            return
        target_columns = config["target_columns"]
        if not set(target_columns).issubset(headers):
            fail("COLUMN_NOT_FOUND", "A selected column no longer exists. Inspect the file again.")
        if config["operation"] == "regex_extract" and config["new_column"].casefold() in {
            x.casefold() for x in headers
        }:
            fail("COLUMN_CONFLICT", "The new column name conflicts with an existing column.")
        stage("TRANSFORMING", 55)
        df = df.repartition(spec["partitions"]).persist(StorageLevel.DISK_ONLY)
        largest = df.select(F.length(F.to_json(F.struct(*[column(name) for name in headers]))).alias("n"))
        if (largest.agg(F.max("n")).first()[0] or 0) > 131072:
            fail("ROW_TOO_LARGE", "A row exceeds the 128 KiB preview safety limit.")
        stats = statistics(df, spec["plan"], config)
        result = transform(df, spec["plan"], config)
        stage("WRITING", 80)
        destination = work / "parquet"
        result.write.mode("errorifexists").option("maxRecordsPerFile", 1000).parquet(str(destination))
        atomic_json(
            work / "output.json",
            {
                **stats,
                "columns": [{"name": name, "type": "string"} for name in result.columns],
                "application_id": spark.sparkContext.applicationId,
                "partitions": df.rdd.getNumPartitions(),
                "executor_count": max(
                    0, int(spark.sparkContext._jsc.sc().getExecutorMemoryStatus().size()) - 1
                ),
                "spark_seconds": round(time.monotonic() - started, 3),
            },
        )
        df.unpersist()
    finally:
        spark.stop()


if __name__ == "__main__":
    spec_path = sys.argv[1]
    try:
        run(spec_path)
    except Exception as exc:
        failure = Path(spec_path).parent / "failure.json"
        if not failure.exists():
            atomic_json(
                failure,
                {
                    "code": "SPARK_EXECUTION_FAILED",
                    "message": "Data processing failed. Check file format, cell limits and worker resources.",
                    "error_type": type(exc).__name__,
                },
            )
        # Do not print exception bodies, Spark configuration, credentials or data rows.
        print("Spark pipeline failed: " + type(exc).__name__, file=sys.stderr)
        sys.exit(1)
